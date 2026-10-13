"""Operations on the Word template: instructions page, markers, placeholders."""

import re

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .errors import TemplateError, UndefinedVariableError
from .variables import NotScalarError, VariableLookupError, lookup_scalar

CONTENT_MARKER = "{{content}}"
END_INSTRUCTIONS_MARKER = "{{end-instructions}}"
RESERVED = {"content", "end-instructions"}
PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][\w.\-]*)\s*\}\}")

W_P = qn("w:p")
W_T = qn("w:t")
W_BR = qn("w:br")
FIELD_TAGS = (qn("w:fldChar"), qn("w:fldSimple"))


def _text_elements(p):
    """``w:t`` elements of paragraph ``p``, excluding paragraphs nested in text boxes."""
    return [t for t in p.iter(W_T) if next(t.iterancestors(W_P)) is p]


def paragraph_text(p):
    return "".join(t.text or "" for t in _text_elements(p))


def _is_page_break_only(el):
    return (
        el.tag == W_P
        and not paragraph_text(el).strip()
        and any(br.get(qn("w:type")) == "page" for br in el.iter(W_BR))
    )


# -- markers -----------------------------------------------------------------------


def instructions_elements(body):
    """Top-level body elements making up the instructions page (empty if no marker).

    That is everything up to and including the ``{{end-instructions}}``
    paragraph, plus a page-break-only paragraph directly after it.
    """
    children = list(body.iterchildren())
    for i, child in enumerate(children):
        if child.tag == W_P and paragraph_text(child).strip() == END_INSTRUCTIONS_MARKER:
            end = i + 1
            if end < len(children) and _is_page_break_only(children[end]):
                end += 1
            return children[:end]
    return []


def strip_instructions(document):
    body = document.element.body
    elements = instructions_elements(body)
    if not elements:
        raise TemplateError(
            f"strip_instructions=True but the template has no paragraph containing only "
            f"{END_INSTRUCTIONS_MARKER!r}; add the marker at the end of the instructions "
            f"page or pass strip_instructions=False"
        )
    for el in elements:
        body.remove(el)


def find_content_anchor(document, skip=()):
    """The paragraph containing only ``{{content}}``, as a python-docx Paragraph."""
    skip = set(skip)
    matches = [
        p
        for child in document.element.body.iterchildren() if child not in skip
        for p in child.iter(W_P) if paragraph_text(p).strip() == CONTENT_MARKER
    ]
    if len(matches) != 1:
        problem = "no" if not matches else f"{len(matches)}"
        raise TemplateError(f"template has {problem} paragraphs containing only "
                            f"{CONTENT_MARKER!r}; expected exactly one")
    return Paragraph(matches[0], document._body)


def remove_anchor(paragraph):
    p = paragraph._p
    if p.pPr is not None and p.pPr.sectPr is not None:
        # The marker paragraph ends a section: keep it (and the section break)
        # but drop its runs.
        for child in list(p):
            if child is not p.pPr:
                p.remove(child)
    else:
        p.getparent().remove(p)


# -- placeholders ------------------------------------------------------------------


def _template_roots(document, skip=()):
    """Body children (minus ``skip``) plus each header and footer root element."""
    skip = set(skip)
    yield from (child for child in document.element.body.iterchildren() if child not in skip)
    parts = {
        rel.target_part
        for rel in document.part.rels.values()
        if not rel.is_external and rel.reltype in (RT.HEADER, RT.FOOTER)
    }
    for part in parts:
        yield part.element


def substitute_placeholders(document, variables, skip=()):
    """Replace ``{{dotted.name}}`` in the body, headers and footers.

    Word often splits text like ``{{version}}`` over several runs, so matching
    is done on the paragraph's combined text and each replacement is written
    into the run where the placeholder starts (keeping that run's formatting).
    """
    undefined = []
    for root in _template_roots(document, skip):
        for p in list(root.iter(W_P)):
            undefined.extend(_substitute_paragraph(p, variables))
    if undefined:
        raise UndefinedVariableError(sorted(set(undefined)))


def _substitute_paragraph(p, variables):
    """Fill placeholders in one paragraph; return the names that could not be filled."""
    texts = _text_elements(p)
    full = "".join(t.text or "" for t in texts)
    matches = [m for m in PLACEHOLDER.finditer(full) if m.group(1) not in RESERVED]
    undefined = []
    for m in reversed(matches):  # back to front so earlier offsets stay valid
        name = m.group(1)
        try:
            value = lookup_scalar(variables, name)
        except NotScalarError as exc:
            undefined.append(f"{name} ({exc.reason})")
        except VariableLookupError:
            undefined.append(name)
        else:
            _replace_range(texts, m.start(), m.end(), value)
    return undefined


def _replace_range(texts, start, end, value):
    """Replace characters ``start:end`` of the joined ``texts`` with ``value``."""
    pos = 0
    for t in texts:
        text = t.text or ""
        t_start, pos = pos, pos + len(text)
        if pos <= start or t_start >= end:
            continue
        a, b = max(start - t_start, 0), min(end - t_start, len(text))
        t.text = text[:a] + (value if t_start <= start else "") + text[b:]
        t.set(qn("xml:space"), "preserve")


# -- settings ----------------------------------------------------------------------

# Elements that must follow w:updateFields in CT_Settings (schema order).
_AFTER_UPDATE_FIELDS = {qn(tag) for tag in (
    "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
    "w:rsids", "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
    "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
    "w:captions", "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary",
    "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator",
)}


def has_fields(document):
    return any(next(root.iter(*FIELD_TAGS), None) is not None for root in _template_roots(document))


def set_update_fields(document):
    """Make Word refresh fields (TOC, page numbers) when the file is opened."""
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        successor = next((c for c in settings.iterchildren() if c.tag in _AFTER_UPDATE_FIELDS), None)
        if successor is not None:
            successor.addprevious(update)
        else:
            settings.append(update)
    update.set(qn("w:val"), "true")
