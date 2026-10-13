"""Walks a docutils doctree and writes python-docx content before an anchor.

Dispatch is by node class through a handler registry, so callers can add
renderers for their own docutils nodes::

    @register_handler(MyNode)
    def render_my_node(renderer, node):
        with renderer.inline_paragraph("Normal"), renderer.formatting(bold=True):
            renderer.render_children(node)

Handlers use the renderer's helpers: ``inline_paragraph``, ``add_text``,
``formatting``, ``render`` and ``render_children``.
"""

import hashlib
import re
from contextlib import contextmanager
from dataclasses import dataclass
from functools import partial
from urllib.parse import urlsplit

from docutils import nodes
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.run import Run

from . import styles as S
from .errors import UnsupportedNodeError
from .extensions import appendix_marker, ref_node
from .numbering import Numbering, set_paragraph_numbering

MONO_FONT = "Courier New"
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
# Anything else (file:, javascript:, UNC paths, ...) could run or open something
# on the reader's machine when clicked.
LINK_SCHEMES = {"http", "https", "mailto"}
RUN_FLAGS = ("bold", "italic", "underline", "strike", "superscript", "subscript")
INLINE_CLASSES = {"underline", "strike"}  # classes set by the :underline:/:strike: roles
W_TBL = qn("w:tbl")

_handlers = {}


def register_handler(*node_classes):
    """Decorator registering ``handler(renderer, node)`` for the given node classes."""
    def decorator(handler):
        for cls in node_classes:
            _handlers[cls] = handler
        return handler
    return decorator


def handler_for(node):
    return next((_handlers[cls] for cls in type(node).__mro__ if cls in _handlers), None)


# -- containers ------------------------------------------------------------------


class AnchorContainer:
    """Inserts block content immediately before the template's {{content}} paragraph."""

    def __init__(self, document, anchor):
        self.document = document
        self.anchor = anchor

    def add_paragraph(self):
        return self.anchor.insert_paragraph_before()

    def add_table(self, rows, cols):
        # Word merges directly adjacent tables into one; keep them apart.
        prev = self.anchor._p.getprevious()
        if prev is not None and prev.tag == W_TBL:
            self.add_paragraph()
        table = self.document.add_table(rows, cols)
        self.anchor._p.addprevious(table._tbl)
        return table

    @property
    def width(self):
        section = self.document.sections[-1]
        return section.page_width - section.left_margin - section.right_margin


class CellContainer:
    """Block content inside a table cell.

    A cell always ends with a paragraph, so an empty trailing paragraph (the
    cell's initial one, or the one python-docx adds after a nested table) is
    reused by the next ``add_paragraph`` instead of leaving a blank line.
    """

    def __init__(self, cell):
        self.cell = cell
        self._spare = cell.paragraphs[0]

    def add_paragraph(self):
        if self._spare is None:
            return self.cell.add_paragraph()
        para, self._spare = self._spare, None
        return para

    def add_table(self, rows, cols):
        spare, self._spare = self._spare, None
        if spare is None and self.cell._tc[-1].tag == W_TBL:
            self.cell.add_paragraph()
        table = self.cell.add_table(rows, cols)
        if spare is not None:
            spare._p.getparent().remove(spare._p)
        self._spare = self.cell.paragraphs[-1]
        return table


# -- renderer --------------------------------------------------------------------


@dataclass
class _ListState:
    style_chain: list
    depth: int
    num_id: int | None = None  # None: the paragraph style supplies numbering
    ilvl: int = 0
    first_paragraph: bool = False


class Renderer:
    def __init__(self, document, anchor):
        self.document = document
        self.styles = S.StyleResolver(document)
        self.numbering = Numbering(document)
        self.containers = [AnchorContainer(document, anchor)]
        self.contexts = []        # "cell" or _ListState, innermost last
        self.paragraph = None     # current Paragraph while rendering inline content
        self.run_parent = None    # element new runs go into (w:p or w:hyperlink)
        self.fmt = {}
        self.section_depth = 0
        self.in_appendix = False  # set by ".. appendix::"
        self.pending_ids = []
        self._bookmarks = {}
        self._next_bookmark_id = 1 + max(
            (int(b.get(qn("w:id"), 0)) for b in document.element.body.iter(qn("w:bookmarkStart"))),
            default=0,
        )

    @property
    def container(self):
        return self.containers[-1]

    @property
    def context(self):
        return self.contexts[-1] if self.contexts else None

    def render(self, node):
        if isinstance(node, nodes.Element) and node.get("ids"):
            self.pending_ids.extend(node["ids"])
            self._emit_bookmarks()
        handler = handler_for(node)
        if handler is None:
            raise UnsupportedNodeError(node)
        handler(self, node)

    def render_children(self, node):
        for child in node.children:
            self.render(child)

    @contextmanager
    def inline_paragraph(self, style_chain=None):
        """Add a paragraph and direct inline content into it until the block exits."""
        para = self.container.add_paragraph()
        if style_chain and (style := self.styles.get(style_chain)) is not None:
            para.style = style
        self.paragraph, self.run_parent = para, para._p
        self._emit_bookmarks()
        try:
            yield para
        finally:
            self.paragraph = self.run_parent = None

    @contextmanager
    def formatting(self, **flags):
        saved = self.fmt
        self.fmt = {**saved, **flags}
        try:
            yield
        finally:
            self.fmt = saved

    @contextmanager
    def in_container(self, container, context="cell"):
        saved = self.paragraph, self.run_parent
        self.containers.append(container)
        self.contexts.append(context)
        self.paragraph = self.run_parent = None
        try:
            yield container
        finally:
            self.containers.pop()
            self.contexts.pop()
            self.paragraph, self.run_parent = saved

    def add_text(self, text):
        if self.paragraph is None:
            raise UnsupportedNodeError(nodes.Text(text), "text outside a paragraph")
        r = OxmlElement("w:r")
        self.run_parent.append(r)
        run = Run(r, self.paragraph)
        run.text = text
        for flag in RUN_FLAGS:
            if self.fmt.get(flag):
                setattr(run.font, flag, True)
        if self.fmt.get("code"):
            self._character_style(run, "HTML Code", font_name=MONO_FONT)
        if self.fmt.get("hyperlink"):
            # Inline code inside a link keeps its code style; colour it like a link.
            style = None if self.fmt.get("code") else "Hyperlink"
            self._character_style(run, style, color=LINK_COLOR, underline=True)
        return run

    def _character_style(self, run, name, font_name=None, color=None, underline=False):
        """Apply character style ``name``, or the given direct formatting if it's missing."""
        style = self.styles.get(name, S.CHARACTER) if name else None
        if style is not None:
            run.style = style
            return
        if font_name:
            run.font.name = font_name
        if color:
            run.font.color.rgb = color
        if underline:
            run.font.underline = True

    # -- bookmarks / hyperlinks ----------------------------------------------

    def bookmark_name(self, node_id):
        """Word bookmark name for a docutils id (hidden: leading underscore, <= 40 chars)."""
        if node_id not in self._bookmarks:
            name = "_r2d_" + re.sub(r"\W", "_", node_id, flags=re.ASCII)
            if len(name) > 40:
                name = f"{name[:31]}_{hashlib.sha256(node_id.encode()).hexdigest()[:8]}"
            self._bookmarks[node_id] = name
        return self._bookmarks[node_id]

    def _emit_bookmarks(self):
        if self.paragraph is None:
            return  # emitted when the next paragraph opens
        for node_id in self.pending_ids:
            bookmark_id = str(self._next_bookmark_id)
            self._next_bookmark_id += 1
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), bookmark_id)
            start.set(qn("w:name"), self.bookmark_name(node_id))
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), bookmark_id)
            self.run_parent.extend([start, end])
        self.pending_ids = []

    @contextmanager
    def hyperlink(self, url=None, anchor_id=None):
        link = OxmlElement("w:hyperlink")
        if url is not None:
            link.set(qn("r:id"), self.paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        else:
            link.set(qn("w:anchor"), self.bookmark_name(anchor_id))
        self.run_parent.append(link)
        saved, self.run_parent = self.run_parent, link
        try:
            with self.formatting(hyperlink=True):
                yield link
        finally:
            self.run_parent = saved


# -- structure -------------------------------------------------------------------


@register_handler(nodes.comment, nodes.target, nodes.substitution_definition,
                  nodes.system_message, nodes.pending)
def _skip(renderer, node):
    pass


@register_handler(nodes.document, nodes.compound, nodes.container)
def _children(renderer, node):
    renderer.render_children(node)


@register_handler(nodes.section)
def _section(renderer, node):
    renderer.section_depth += 1
    try:
        renderer.render_children(node)
    finally:
        renderer.section_depth -= 1


@register_handler(appendix_marker)
def _appendix_marker(renderer, node):
    renderer.in_appendix = True


@register_handler(nodes.title)
def _title(renderer, node):
    if isinstance(node.parent, nodes.section):
        chain = (S.appendix_chain if renderer.in_appendix else S.heading_chain)(renderer.section_depth)
    elif isinstance(node.parent, nodes.table):
        chain = ["Caption", "Normal"]
    else:
        raise UnsupportedNodeError(node, f"title inside <{node.parent.tagname}>")
    with renderer.inline_paragraph(chain):
        renderer.render_children(node)


@register_handler(nodes.paragraph)
def _paragraph(renderer, node):
    state = renderer.context
    if not isinstance(state, _ListState):
        style_chain = None
    elif state.first_paragraph:
        state.first_paragraph = False
        with renderer.inline_paragraph(state.style_chain) as para:
            if state.num_id is not None:
                set_paragraph_numbering(para, state.num_id, state.ilvl)
            renderer.render_children(node)
        return
    else:
        style_chain = S.list_continue_chain(state.depth)
    with renderer.inline_paragraph(style_chain):
        renderer.render_children(node)


# -- lists -----------------------------------------------------------------------


def _list(renderer, node, kind):
    depth = 1 + sum(isinstance(c, _ListState) for c in renderer.contexts)
    chain = S.list_chain(kind, depth)
    state = _ListState(chain, depth)
    style_numbering = renderer.numbering.for_style(renderer.styles.get(chain))
    enumtype = node.get("enumtype", "arabic")
    start = node.get("start", 1)

    if kind == "enum" and enumtype == "arabic" and style_numbering is not None:
        # Restart the style's own numbering so each RST list counts from its start.
        num_id, state.ilvl = style_numbering
        state.num_id = renderer.numbering.restart(num_id, state.ilvl, start)
    elif kind == "enum" or style_numbering is None:
        # Letters/roman numerals, or a style without numbering: generate a definition.
        state.ilvl = depth - 1
        state.num_id = renderer.numbering.generated(
            enumtype if kind == "enum" else "bullet", state.ilvl, start)

    renderer.contexts.append(state)
    try:
        renderer.render_children(node)
    finally:
        renderer.contexts.pop()


register_handler(nodes.bullet_list)(partial(_list, kind="bullet"))
register_handler(nodes.enumerated_list)(partial(_list, kind="enum"))


@register_handler(nodes.list_item)
def _list_item(renderer, node):
    state = renderer.context
    state.first_paragraph = True
    if not node.children or not isinstance(node[0], nodes.paragraph):
        _paragraph(renderer, nodes.paragraph())  # still emit the bullet/number
    renderer.render_children(node)
    state.first_paragraph = False


# -- blocks ----------------------------------------------------------------------


@register_handler(nodes.literal_block, nodes.doctest_block)
def _literal_block(renderer, node):
    style = renderer.styles.get([*S.MONOSPACE_PARAGRAPH_STYLES, "Normal"])
    with renderer.inline_paragraph() as para:
        if style is not None:
            para.style = style
        run = renderer.add_text(node.astext())  # "\n" becomes a line break
        if style is None or style.name not in S.MONOSPACE_PARAGRAPH_STYLES:
            run.font.name = MONO_FONT
            run.font.size = Pt(9)


@register_handler(nodes.block_quote)
def _block_quote(renderer, node):
    for child in node.children:
        if isinstance(child, (nodes.paragraph, nodes.attribution)):
            with renderer.inline_paragraph(["Quote", "Normal"]):
                if isinstance(child, nodes.attribution):
                    renderer.add_text("— ")
                renderer.render_children(child)
        else:
            renderer.render(child)


@register_handler(nodes.admonition, nodes.attention, nodes.caution, nodes.danger,
                  nodes.error, nodes.hint, nodes.important, nodes.note, nodes.tip,
                  nodes.warning)
def _admonition(renderer, node):
    """A boxed single-cell table with a bold title line."""
    # Only the generic ``.. admonition::`` carries a title node, as its first child.
    title = node[0] if isinstance(node, nodes.admonition) else None
    table = renderer.container.add_table(1, 1)
    if (style := renderer.styles.get("Table Grid", S.TABLE)) is not None:
        table.style = style
    with renderer.in_container(CellContainer(table.cell(0, 0))):
        with renderer.inline_paragraph(), renderer.formatting(bold=True):
            if title is not None:
                renderer.render_children(title)
            else:
                renderer.add_text(node.tagname.capitalize())
        for child in node.children[1 if title is not None else 0:]:
            renderer.render(child)


# -- tables ------------------------------------------------------------------------


def _grid_layout(rows):
    """Yield ``(row, col, morerows, morecols, entry)`` with spans placed on a grid."""
    occupied = set()
    for r, row in enumerate(rows):
        c = 0
        for entry in row.children:
            while (r, c) in occupied:
                c += 1
            mr, mc = entry.get("morerows", 0), entry.get("morecols", 0)
            occupied.update((rr, cc) for rr in range(r, r + mr + 1) for cc in range(c, c + mc + 1))
            yield r, c, mr, mc, entry
            c += mc + 1


def _children_of_type(node, cls):
    return [child for child in node.children if isinstance(child, cls)]


def _first_child_of_type(node, cls):
    return next((child for child in node.children if isinstance(child, cls)), None)


@register_handler(nodes.table)
def _table(renderer, node):
    title = _first_child_of_type(node, nodes.title)
    tgroup = _first_child_of_type(node, nodes.tgroup)
    if tgroup is None:
        raise UnsupportedNodeError(node, "table without tgroup")
    colwidths = [c.get("colwidth", 1) for c in _children_of_type(tgroup, nodes.colspec)]
    head_rows = [row for part in _children_of_type(tgroup, nodes.thead) for row in part.children]
    body_rows = [row for part in _children_of_type(tgroup, nodes.tbody) for row in part.children]
    rows = head_rows + body_rows
    if not rows:
        raise UnsupportedNodeError(node, "table with no rows")
    layout = list(_grid_layout(rows))
    ncols = tgroup["cols"]

    if title is not None:
        renderer.render(title)
    table = renderer.container.add_table(len(rows), ncols)
    if (style := renderer.styles.get("Table Grid", S.TABLE)) is not None:
        table.style = style

    for r, c, mr, mc, _ in layout:
        if mr or mc:
            table.cell(r, c).merge(table.cell(r + mr, c + mc))

    if len(colwidths) == ncols and isinstance(renderer.container, AnchorContainer):
        table.autofit = False
        total = sum(colwidths)
        widths = [int(renderer.container.width * w / total) for w in colwidths]
        for column, width in zip(table.columns, widths):
            column.width = width
        for r, c, _, mc, _ in layout:
            table.cell(r, c).width = sum(widths[c:c + mc + 1])

    for row in table.rows[:len(head_rows)]:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # repeat on each page

    for r, c, _, _, entry in layout:
        with renderer.in_container(CellContainer(table.cell(r, c))), \
                renderer.formatting(bold=r < len(head_rows)):
            renderer.render_children(entry)


# -- inline ------------------------------------------------------------------------


@register_handler(nodes.Text)
def _text(renderer, node):
    if text := node.astext().replace("\n", " "):
        renderer.add_text(text)


def _formatted(**flags):
    def handler(renderer, node):
        with renderer.formatting(**flags):
            renderer.render_children(node)
    return handler


register_handler(nodes.strong)(_formatted(bold=True))
register_handler(nodes.emphasis, nodes.title_reference)(_formatted(italic=True))
register_handler(nodes.literal)(_formatted(code=True))
register_handler(nodes.superscript)(_formatted(superscript=True))
register_handler(nodes.subscript)(_formatted(subscript=True))


@register_handler(nodes.inline)
def _inline(renderer, node):
    flags = {cls: True for cls in node.get("classes", []) if cls in INLINE_CLASSES}
    with renderer.formatting(**flags):
        renderer.render_children(node)


@register_handler(nodes.reference, ref_node)
def _reference(renderer, node):
    # ref_node (:ref:) gets its refid from api._resolve_refs before rendering.
    if "refuri" in node:
        if urlsplit(node["refuri"]).scheme.lower() not in LINK_SCHEMES:
            raise UnsupportedNodeError(
                node, f"link to {node['refuri']!r}: only {', '.join(sorted(LINK_SCHEMES))} links are allowed")
        link = renderer.hyperlink(url=node["refuri"])
    elif "refid" in node:
        link = renderer.hyperlink(anchor_id=node["refid"])
    else:
        raise UnsupportedNodeError(node, "unresolved reference")
    with link:
        renderer.render_children(node)
