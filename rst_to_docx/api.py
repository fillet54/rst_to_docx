import io
import os
import warnings

import docx
from docutils import nodes
from docutils.core import publish_doctree
from docx.document import Document as DocumentType

from .errors import RstError, RstWarning
from .extensions import SETTING_JIRA_URL, SETTING_VARIABLES, ref_node
from .renderer import Renderer
from .template import (
    find_content_anchor, has_fields, instructions_elements, remove_anchor,
    set_update_fields, strip_instructions as _strip_instructions, substitute_placeholders,
)


def parse(source, *, variables=None, jira_url=None, source_path=None):
    """Parse RST into a docutils doctree, raising RstError on any error-level problem."""
    settings = {
        "doctitle_xform": False,
        "sectsubtitle_xform": False,
        "docinfo_xform": False,
        "report_level": 2,
        "halt_level": 5,
        "warning_stream": io.StringIO(),
        "syntax_highlight": "none",
        "smart_quotes": True,
        "raw_enabled": False,
        "input_encoding": "utf-8",
        SETTING_VARIABLES: variables or {},
        SETTING_JIRA_URL: jira_url,
    }
    doctree = publish_doctree(source, source_path=source_path, settings_overrides=settings)

    errors = []
    for msg in doctree.findall(nodes.system_message):
        text = "%s:%s: %s" % (msg.get("source") or source_path or "<string>",
                              msg.get("line", "?"), msg.children[0].astext() if msg.children else "")
        if msg["level"] >= 3:
            errors.append(text)
        elif msg["level"] == 2:
            warnings.warn(text, RstWarning, stacklevel=3)
    errors.extend(_resolve_refs(doctree, source_path))
    if errors:
        raise RstError(errors)
    return doctree


def _resolve_refs(doctree, source_path):
    errors = []
    for node in list(doctree.findall(ref_node)):
        label = node["reflabel"]
        refid = doctree.nameids.get(nodes.fully_normalize_name(label))
        target = doctree.ids.get(refid) if refid else None
        if target is None:
            errors.append("%s:%s: :ref: target %r not found"
                          % (node.source or source_path or "<string>", node.line, label))
            continue
        node["refid"] = refid
        if not node["explicit_title"]:
            title = target.next_node(nodes.title) if isinstance(target, (nodes.section, nodes.table)) else None
            text = title.astext() if title is not None else label
            node.children = []
            node += nodes.Text(text)
    return errors


def render(source, template, output=None, *, variables=None, jira_url=None,
           strip_instructions=True, update_fields=None, source_path=None):
    """Render reStructuredText into a Word template.

    :param source: RST text, or a path-like object pointing at an .rst file.
    :param template: path to the .docx template, a file-like object, or a
        python-docx Document (which is modified in place).
    :param output: optional path or file-like object to save the result to.
    :param variables: dict used by ``:var:``, ``var-table`` and template
        ``{{placeholders}}``. Dotted paths walk nested dicts/lists/attributes.
    :param jira_url: base URL for ``:jira:`` links, e.g.
        ``"https://jira.example.com/browse/"`` (or containing ``{key}``).
    :param strip_instructions: remove everything up to and including the
        ``{{end-instructions}}`` paragraph.
    :param update_fields: ask Word to refresh fields (TOC, page refs) on open.
        ``None`` (default) does so only when the template contains fields.
    :param source_path: file name used in error messages and for resolving
        relative paths (``include``, ``csv-table :file:``) when ``source`` is text.
    :returns: the python-docx Document.
    """
    if isinstance(source, os.PathLike):
        source_path = source_path or os.fspath(source)
        with open(source, encoding="utf-8") as fh:
            source = fh.read()
    variables = variables or {}

    doctree = parse(source, variables=variables, jira_url=jira_url, source_path=source_path)

    document = template if isinstance(template, DocumentType) else docx.Document(template)
    body = document.element.body
    if strip_instructions:
        _strip_instructions(document)
        skip = []
    else:
        skip = instructions_elements(body)

    anchor = find_content_anchor(document, skip=skip)
    # Placeholders are filled before the RST goes in, so RST text containing
    # "{{...}}" is never touched.
    substitute_placeholders(document, variables, skip=skip)

    Renderer(document, anchor, doctree).render(doctree)
    remove_anchor(anchor)

    if update_fields is None:
        update_fields = has_fields(document)
    if update_fields:
        set_update_fields(document)

    if output is not None:
        document.save(output)
    return document
