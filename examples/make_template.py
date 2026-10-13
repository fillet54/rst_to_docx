"""Build examples/template.docx, a clean SVD template.

Look:
  * only the styles the renderer and the template itself use (see STYLES)
  * all text black Times New Roman; code is Courier New
  * Heading 1-4 are 12 pt bold and numbered 1., 1.1, 1.1.1, 1.1.1.1
  * every page, title page included, carries the document number, revision
    letter and release date in the header and footer, and the classification
    in red at the bottom

Layout:
  page 1  instructions page: a sample of every style, ends with {{end-instructions}}
  page 2  title page with {{placeholders}}
  page 3  table of contents field
  then    {{content}}

In practice you'd design the template in Word; this script just produces a
reproducible example.
"""

import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, RGBColor

HEADING_LEVELS = 4
LIST_LEVELS = 3
FONT = "Times New Roman"
MONO_FONT = "Courier New"
BODY_SIZE = 24  # half-points
HEADING_SIZE = 24
RED = RGBColor(0xFF, 0x00, 0x00)
TEXT_WIDTH = 9360  # twips: Letter paper with 1 inch margins

# numbering.xml ids: one multilevel list for the headings, one list per list style.
HEADING_NUM = 1
BULLET_NUM = 2  # .. BULLET_NUM + LIST_LEVELS - 1
NUMBER_NUM = BULLET_NUM + LIST_LEVELS


def _fonts(name):
    return f'<w:rFonts w:ascii="{name}" w:hAnsi="{name}" w:eastAsia="{name}" w:cs="{name}"/>'


def _size(half_points):
    return f'<w:sz w:val="{half_points}"/><w:szCs w:val="{half_points}"/>'


_LOWER_CASE_NAMES = ("Heading ", "TOC ", "Header", "Footer", "Caption")


def _style(name, kind="paragraph", based_on="Normal", next_style=None, ppr="", rpr="", extra="",
           default=False, style_id=None):
    style_id = style_id or name.replace(" ", "")
    default_attr = ' w:default="1"' if default else ""
    # Word stores these built-in names in lower case ("heading 1", "toc 1", "header").
    stored = name.lower() if name.startswith(_LOWER_CASE_NAMES) and name != "TOC Heading" else name
    parts = [f'<w:style w:type="{kind}" w:styleId="{style_id}"{default_attr}>',
             f'<w:name w:val="{stored}"/>']
    if based_on and not default:
        parts.append(f'<w:basedOn w:val="{based_on}"/>')
    if next_style:
        parts.append(f'<w:next w:val="{next_style}"/>')
    if kind == "paragraph":
        parts.append("<w:qFormat/>")
    if ppr:
        parts.append(f"<w:pPr>{ppr}</w:pPr>")
    if rpr:
        parts.append(f"<w:rPr>{rpr}</w:rPr>")
    parts.append(extra + "</w:style>")
    return "".join(parts)


def _heading_styles():
    for i in range(HEADING_LEVELS):
        yield _style(
            f"Heading {i + 1}", next_style="Normal",
            ppr=f'<w:keepNext/><w:keepLines/><w:numPr><w:ilvl w:val="{i}"/>'
                f'<w:numId w:val="{HEADING_NUM}"/></w:numPr>'
                f'<w:spacing w:before="240" w:after="120"/><w:outlineLvl w:val="{i}"/>',
            rpr=f"{_fonts(FONT)}<w:b/><w:bCs/>{_size(HEADING_SIZE)}",
        )


def _toc_styles():
    for i in range(3):
        yield _style(
            f"TOC {i + 1}", next_style="Normal",
            ppr=f'<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{TEXT_WIDTH}"/></w:tabs>'
                f'<w:spacing w:after="60"/><w:ind w:left="{240 * i}"/>',
        )


def _list_styles():
    for i in range(LIST_LEVELS):
        n = f" {i + 1}" if i else ""
        for kind, num_id in (("List Bullet", BULLET_NUM + i), ("List Number", NUMBER_NUM + i)):
            yield _style(kind + n, ppr=f'<w:numPr><w:numId w:val="{num_id}"/></w:numPr>'
                                       f"<w:contextualSpacing/>")
        yield _style("List Continue" + n,
                     ppr=f'<w:ind w:left="{360 * (i + 1)}"/><w:contextualSpacing/>')


_TABLE_MARGINS = (
    '<w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/>'
    '<w:left w:w="108" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/>'
    '<w:right w:w="108" w:type="dxa"/></w:tblCellMar>'
)
_HEADER_TABS = (
    f'<w:tabs><w:tab w:val="center" w:pos="{TEXT_WIDTH // 2}"/>'
    f'<w:tab w:val="right" w:pos="{TEXT_WIDTH}"/></w:tabs><w:spacing w:after="0"/>'
)

DOC_DEFAULTS = (
    f"<w:docDefaults><w:rPrDefault><w:rPr>{_fonts(FONT)}"
    f'<w:color w:val="000000"/>{_size(BODY_SIZE)}'
    f'<w:lang w:val="en-US" w:eastAsia="en-US" w:bidi="ar-SA"/></w:rPr></w:rPrDefault>'
    f'<w:pPrDefault><w:pPr><w:spacing w:after="120"/></w:pPr></w:pPrDefault></w:docDefaults>'
)

# Every style in the template. The first four are the defaults Word expects.
STYLES = [
    _style("Normal", default=True),
    _style("Default Paragraph Font", kind="character", default=True),
    _style("Normal Table", kind="table", default=True, style_id="TableNormal",
           extra=f"<w:tblPr>{_TABLE_MARGINS}</w:tblPr>"),
    _style("No List", kind="numbering", default=True),
    *_heading_styles(),
    _style("Title", next_style="Normal",
           ppr='<w:spacing w:after="240"/><w:jc w:val="center"/>', rpr=f"<w:b/><w:bCs/>{_size(40)}"),
    _style("Subtitle", next_style="Normal",
           ppr='<w:spacing w:after="480"/><w:jc w:val="center"/>', rpr=_size(28)),
    _style("TOC Heading", next_style="Normal", ppr='<w:keepNext/><w:spacing w:after="240"/>',
           rpr=f"<w:b/><w:bCs/>{_size(HEADING_SIZE)}"),
    *_toc_styles(),
    *_list_styles(),
    _style("Plain Text", rpr=f"{_fonts(MONO_FONT)}{_size(18)}"),
    _style("Caption", next_style="Normal", ppr='<w:keepNext/><w:spacing w:before="120"/>',
           rpr="<w:b/><w:bCs/>"),
    _style("Quote", ppr='<w:ind w:left="720" w:right="720"/>', rpr="<w:i/><w:iCs/>"),
    _style("Header", ppr=_HEADER_TABS, rpr=_size(20)),
    _style("Footer", ppr=_HEADER_TABS, rpr=_size(20)),
    _style("Hyperlink", kind="character", based_on="DefaultParagraphFont",
           rpr='<w:u w:val="single"/>'),
    _style("HTML Code", kind="character", based_on="DefaultParagraphFont",
           rpr=f"{_fonts(MONO_FONT)}{_size(20)}"),
    _style("Table Grid", kind="table", based_on="TableNormal",
           ppr='<w:spacing w:after="0"/>',
           extra="<w:tblPr><w:tblBorders>" + "".join(
               f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
               for side in ("top", "left", "bottom", "right", "insideH", "insideV")
           ) + f"</w:tblBorders>{_TABLE_MARGINS}</w:tblPr>"),
]


def set_styles(doc):
    """Replace the python-docx default styles (160 or so) with STYLES."""
    styles = doc.styles.element
    latent = styles.find(qn("w:latentStyles"))  # how Word lists the built-ins we left out
    new = parse_xml(f'<w:styles {nsdecls("w")}>{DOC_DEFAULTS}{"".join(STYLES)}</w:styles>')
    for child in list(styles):
        styles.remove(child)
    styles.extend(list(new))
    styles.insert(1, latent)

    # Drop the Word 2010 copy of the old styles and the default template's thumbnail.
    for rels, suffix in ((doc.part.rels, "/stylesWithEffects"), (doc.part.package.rels, "/thumbnail")):
        for r_id in [r_id for r_id, rel in rels.items() if rel.reltype.endswith(suffix)]:
            del rels[r_id]

    # Make Times New Roman the theme's heading and body font as well.
    theme = doc.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    theme._blob = re.sub(rb'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*',
                         rb"\g<1>" + FONT.encode(), theme.blob)


def _level(ilvl, fmt, text, left, hanging=None, style=None, font=None):
    hanging = left if hanging is None else hanging
    return (
        f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
        + (f'<w:pStyle w:val="{style}"/>' if style else "")
        + f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr>'
        + (f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:hint="default"/></w:rPr>'
           if font else "")
        + "</w:lvl>"
    )


def set_numbering(doc):
    """Replace numbering.xml: heading numbers plus one list per list style."""
    # Level 1 is "1.", deeper levels are "1.1", "1.1.1", ...
    heading_levels = [
        _level(i, "decimal", ".".join(f"%{n + 1}" for n in range(i + 1)) + ("" if i else "."),
               left=432 + 144 * i, style=f"Heading{i + 1}")
        for i in range(HEADING_LEVELS)
    ]
    abstracts = [("multilevel", "".join(heading_levels))]
    bullets = (("", "Symbol"), ("o", MONO_FONT), ("", "Wingdings"))
    for i in range(LIST_LEVELS):
        n = str(i + 1) if i else ""
        char, font = bullets[i % len(bullets)]
        abstracts.append(("singleLevel", _level(0, "bullet", char, 360 * (i + 1), 360,
                                                 f"ListBullet{n}", font)))
    for i in range(LIST_LEVELS):
        n = str(i + 1) if i else ""
        abstracts.append(("singleLevel", _level(0, "decimal", "%1.", 360 * (i + 1), 360,
                                                 f"ListNumber{n}")))
    # Order above matches HEADING_NUM, BULLET_NUM.., NUMBER_NUM..
    xml = "".join(
        f'<w:abstractNum w:abstractNumId="{i}"><w:multiLevelType w:val="{kind}"/>{levels}'
        f"</w:abstractNum>" for i, (kind, levels) in enumerate(abstracts, 1)
    ) + "".join(
        f'<w:num w:numId="{i}"><w:abstractNumId w:val="{i}"/></w:num>'
        for i in range(1, len(abstracts) + 1)
    )
    numbering = doc.part.numbering_part.element
    for child in list(numbering):
        numbering.remove(child)
    numbering.extend(list(parse_xml(f'<w:numbering {nsdecls("w")}>{xml}</w:numbering>')))


def add_field(paragraph, instr, placeholder=""):
    """Append a complex field (e.g. TOC, PAGE) showing ``placeholder`` until updated."""
    def field_run(child):
        paragraph._p.append(parse_xml(f"<w:r {nsdecls('w')}>{child}</w:r>"))

    field_run('<w:fldChar w:fldCharType="begin" w:dirty="true"/>')
    field_run(f'<w:instrText xml:space="preserve"> {instr} </w:instrText>')
    field_run('<w:fldChar w:fldCharType="separate"/>')
    paragraph.add_run(placeholder)
    field_run('<w:fldChar w:fldCharType="end"/>')


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def instructions_page(doc):
    doc.add_paragraph("Template instructions (removed when rendering)", style="Title")
    doc.add_paragraph(
        "Everything up to and including the {{end-instructions}} paragraph below is "
        "deleted by rst_to_docx. This page shows one example of every style the "
        "renderer uses. If you add a built-in style in Word, use it once here: Word "
        "only saves a built-in style into the file once it has been used."
    )
    doc.add_paragraph(
        "Markers: put {{content}} alone in a paragraph where the RST body goes. "
        "Anywhere else (title page, headers, footers, tables) {{dotted.name}} is "
        "replaced from the variables dict."
    )
    for level in range(1, HEADING_LEVELS + 1):
        doc.add_paragraph(f"Heading {level} sample", style=f"Heading {level}")
    doc.add_paragraph("Normal paragraph sample.", style="Normal")
    for kind in ("List Bullet", "List Number", "List Continue"):
        for i in range(LIST_LEVELS):
            name = kind + (f" {i + 1}" if i else "")
            doc.add_paragraph(f"{name} sample", style=name)
    doc.add_paragraph("Plain Text (literal block) sample", style="Plain Text")
    doc.add_paragraph("Caption (table title) sample", style="Caption")
    doc.add_paragraph("Quote sample", style="Quote")
    p = doc.add_paragraph("Character styles: ")
    p.add_run("Hyperlink", style="Hyperlink")
    p.add_run(" and ")
    p.add_run("HTML Code", style="HTML Code")
    table = doc.add_table(rows=2, cols=2, style="Table Grid")
    table.cell(0, 0).text = "Table Grid"
    table.cell(0, 1).text = "sample"
    doc.add_paragraph("{{end-instructions}}")
    page_break(doc)


def title_page(doc):
    for _ in range(8):
        doc.add_paragraph()
    doc.add_paragraph("{{project}}", style="Title")
    doc.add_paragraph("Software Version Description", style="Subtitle")
    rows = [
        ("Document number", "{{doc.number}}"),
        ("Revision", "{{doc.revision}}"),
        ("Software version", "{{release.version}}"),
        ("Release date", "{{release.date}}"),
        ("Prepared by", "{{doc.author}}"),
    ]
    table = doc.add_table(rows=len(rows), cols=2, style="Table Grid")
    for (label, value), row in zip(rows, table.rows):
        row.cells[0].text = label
        row.cells[1].text = value
    page_break(doc)


def toc_page(doc):
    doc.add_paragraph("Table of Contents", style="TOC Heading")
    add_field(doc.add_paragraph(), r'TOC \o "1-3" \h \z \u',
              "Right-click and choose Update Field to build the table of contents.")
    page_break(doc)


def header_footer(doc):
    """Document number, revision and release date on every page, title page included."""
    section = doc.sections[0]
    section.header.paragraphs[0].text = (
        "{{doc.number}}\tRevision {{doc.revision}}\t{{release.date}}"
    )
    footer = section.footer.paragraphs[0]
    footer.text = "{{doc.number}} Rev {{doc.revision}}\t{{release.date}}\tPage "
    add_field(footer, "PAGE", "1")
    footer.add_run(" of ")
    add_field(footer, "NUMPAGES", "1")
    marking = section.footer.add_paragraph(style="Footer")
    marking.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = marking.add_run("{{doc.classification}}")
    run.bold = True
    run.font.color.rgb = RED


def build(path):
    doc = docx.Document()
    section = doc.sections[0]  # margins first: new tables are sized to the text width
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)
    set_styles(doc)
    set_numbering(doc)
    instructions_page(doc)
    title_page(doc)
    toc_page(doc)
    doc.add_paragraph("{{content}}")
    header_footer(doc)
    doc.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("template.docx")
    print(build(out))
