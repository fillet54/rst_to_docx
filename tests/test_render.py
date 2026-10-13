import io
import sys
import warnings
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from rst_to_docx import (
    MissingStyleWarning, RstError, TemplateError, UndefinedVariableError,
    UnsupportedNodeError, render,
)

sys.path.insert(0, str(Path(__file__).parents[1] / "examples"))
from make_template import build  # noqa: E402

VARS = {
    "project": "Widget",
    "release": {"version": "1.4.2", "date": "2026-10-13"},
    "doc": {"number": "SVD-1", "revision": "B", "classification": "UNCLASSIFIED", "author": "QA"},
}
JIRA = "https://jira.example.com/browse/"


@pytest.fixture(scope="session")
def template(tmp_path_factory):
    return build(tmp_path_factory.mktemp("tpl") / "template.docx")


def simple_template(body_paragraphs, header=None):
    doc = docx.Document()
    for text in body_paragraphs:
        doc.add_paragraph(text)
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def do_render(rst, template, **kw):
    kw.setdefault("variables", VARS)
    kw.setdefault("jira_url", JIRA)
    buf = io.BytesIO()
    render(rst, template, buf, **kw)
    buf.seek(0)
    return docx.Document(buf)  # round-trip through a saved file


def texts(doc):
    # paragraph.text skips hyperlink runs in python-docx < 1.0; read all w:t instead.
    return ["".join(t.text or "" for t in p._p.iter(qn("w:t"))) for p in doc.paragraphs]


def body_xml(doc):
    return doc.element.body.xml


# -- template handling -----------------------------------------------------------


def test_instructions_stripped_by_default(template):
    doc = do_render("Hello", template)
    t = texts(doc)
    assert not any("Template instructions" in x for x in t)
    assert "{{end-instructions}}" not in t
    assert "{{content}}" not in t
    assert "Hello" in t


def test_instructions_kept_when_requested(template):
    doc = do_render("Hello", template, strip_instructions=False)
    t = texts(doc)
    assert any("Template instructions" in x for x in t)
    # placeholders on the instructions page are left alone
    assert any("{{dotted.name}}" in x for x in t)


def test_missing_end_marker_raises():
    with pytest.raises(TemplateError, match="end-instructions"):
        render("x", simple_template(["{{content}}"]))


def test_no_end_marker_ok_when_not_stripping():
    doc = do_render("x", simple_template(["{{content}}"]), strip_instructions=False)
    assert texts(doc) == ["x"]


def test_missing_content_marker_raises():
    with pytest.raises(TemplateError, match="content"):
        render("x", simple_template(["nothing"]), strip_instructions=False)


def test_content_inserted_at_marker_position():
    tpl = simple_template(["{{end-instructions}}", "before", "{{content}}", "after"])
    assert texts(do_render("middle", tpl)) == ["before", "middle", "after"]


def test_placeholders_in_cover_table_and_header(template):
    doc = do_render("x", template)
    cover = doc.tables[0]
    assert cover.cell(2, 1).text == "1.4.2"
    assert "Widget" in texts(doc)
    section = doc.sections[0]
    assert not section.different_first_page_header_footer  # the title page gets them too
    assert section.header.paragraphs[0].text == "SVD-1\tRevision B\t2026-10-13"
    info, marking = section.footer.paragraphs
    assert info.text.startswith("SVD-1 Rev B\t2026-10-13\tPage ")
    assert marking.text == "UNCLASSIFIED"
    assert str(marking.runs[0].font.color.rgb) == "FF0000"


def test_template_styles_are_minimal_black_and_numbered(template):
    doc = docx.Document(template)
    used = {p.style.name for p in doc.paragraphs} | {"Table Grid", "Hyperlink", "HTML Code"}
    builtin_defaults = {"Normal", "Default Paragraph Font", "Normal Table", "No List"}
    toc_and_page = {"toc 1", "toc 2", "toc 3", "Header", "Footer"}
    assert {s.name for s in doc.styles} == used | builtin_defaults | toc_and_page
    # Black comes from the document defaults; no style overrides it.
    styles = doc.styles.element
    assert [c.get(qn("w:val")) for c in styles.iter(qn("w:color"))] == ["000000"]
    numbering = doc.part.numbering_part.element
    for level, text in enumerate(["%1.", "%1.%2", "%1.%2.%3"]):
        heading = doc.styles[f"Heading {level + 1}"]
        assert heading.font.name == "Times New Roman" and heading.font.size.pt == 12
        num_pr = heading.element.pPr.numPr
        abstract_id = numbering.num_having_numId(num_pr.numId.val).abstractNumId.val
        lvl = numbering.xpath(
            f"w:abstractNum[@w:abstractNumId={abstract_id}]/w:lvl[@w:ilvl={num_pr.ilvl.val}]")[0]
        assert lvl.find(qn("w:lvlText")).get(qn("w:val")) == text


def test_placeholder_split_across_runs():
    d = docx.Document()
    d.add_paragraph("{{end-instructions}}")
    p = d.add_paragraph()
    p.add_run("v{{rel")
    p.add_run("ease.ver").bold = True
    p.add_run("sion}} done")
    d.add_paragraph("{{content}}")
    doc = do_render("x", d)
    assert doc.paragraphs[0].text == "v1.4.2 done"


def test_undefined_placeholder_raises():
    tpl = simple_template(["{{end-instructions}}", "{{nope}} {{release.bad}}", "{{content}}"])
    with pytest.raises(UndefinedVariableError) as exc:
        render("x", tpl, variables=VARS)
    assert exc.value.names == ["nope", "release.bad"]


def test_rst_text_with_braces_not_substituted():
    doc = do_render("literal ``{{project}}`` stays", simple_template(["{{content}}"]),
                    strip_instructions=False)
    assert "{{project}}" in texts(doc)[0]


def test_update_fields_set_when_template_has_toc(template):
    doc = do_render("x", template)
    assert doc.settings.element.find(qn("w:updateFields")) is not None
    plain = do_render("x", simple_template(["{{content}}"]), strip_instructions=False)
    assert plain.settings.element.find(qn("w:updateFields")) is None


# -- variables ---------------------------------------------------------------------


def test_var_role():
    doc = do_render("Version :var:`release.version` of :var:`project`.",
                    simple_template(["{{content}}"]), strip_instructions=False)
    assert texts(doc) == ["Version 1.4.2 of Widget."]


def test_var_lists_and_attributes():
    class Obj:
        name = "attr"
    doc = do_render(":var:`items.1` :var:`obj.name`", simple_template(["{{content}}"]),
                    strip_instructions=False, variables={"items": ["a", "b"], "obj": Obj()})
    assert texts(doc) == ["b attr"]


def test_undefined_vars_reported_together_with_lines():
    rst = "One :var:`missing`.\n\nTwo :var:`release.nope`.\n"
    with pytest.raises(RstError) as exc:
        render(rst, simple_template(["{{content}}"]), variables=VARS, strip_instructions=False)
    msgs = exc.value.messages
    assert len(msgs) == 2
    assert ":1:" in msgs[0] and "missing" in msgs[0]
    assert ":3:" in msgs[1] and "release.nope" in msgs[1]


def test_non_scalar_var_rejected():
    with pytest.raises(RstError, match="not a scalar"):
        render(":var:`release`", simple_template(["{{content}}"]), variables=VARS,
               strip_instructions=False)


# -- structure and inline ----------------------------------------------------------


def test_headings_start_at_level_one(template):
    rst = "Top\n===\n\ntext\n\nSub\n---\n\nmore\n"
    doc = do_render(rst, template)
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Top"] == "Heading 1"
    assert styles["Sub"] == "Heading 2"


def test_inline_formatting():
    doc = do_render("**b** *i* ``c`` :underline:`u` :strike:`s` H\\ :sub:`2`\\ O",
                    simple_template(["{{content}}"]), strip_instructions=False)
    runs = {r.text: r for r in doc.paragraphs[0].runs}
    assert runs["b"].bold
    assert runs["i"].italic
    assert runs["c"].font.name == "Courier New"  # no HTML Code style in this template
    assert runs["u"].underline
    assert runs["s"].font.strike
    assert runs["2"].font.subscript


def test_code_uses_character_style_when_available(template):
    doc = do_render("``x``", template)
    run = next(r for p in doc.paragraphs for r in p.runs if r.text == "x")
    assert run.style.name == "HTML Code"


def test_smart_dashes():
    doc = do_render("a -- b", simple_template(["{{content}}"]), strip_instructions=False)
    assert texts(doc) == ["a – b"]


def test_literal_block_keeps_lines(template):
    doc = do_render("Code::\n\n   line 1\n     line 2\n", template)
    p = next(p for p in doc.paragraphs if p.style.name == "Plain Text")
    assert p.text == "line 1\n  line 2"


def _num_id(p):
    numPr = p._p.pPr.numPr if p._p.pPr is not None else None
    return numPr.numId.val if numPr is not None and numPr.numId is not None else None


def test_lists_nested_and_numbered_restart(template):
    rst = """
* a
* b

  - b1

#. one
#. two

text

#. again
"""
    doc = do_render(rst, template)
    by_text = {p.text: p for p in doc.paragraphs}
    assert by_text["a"].style.name == "List Bullet"
    assert by_text["b1"].style.name == "List Bullet 2"
    assert by_text["one"].style.name == "List Number"
    assert _num_id(by_text["one"]) == _num_id(by_text["two"])
    assert _num_id(by_text["again"]) != _num_id(by_text["one"])


def test_list_item_continuation_paragraph(template):
    doc = do_render("* first\n\n  second para\n", template)
    by_text = {p.text: p for p in doc.paragraphs}
    assert by_text["second para"].style.name == "List Continue"


def test_lists_fall_back_to_generated_numbering():
    d = docx.Document()
    for name in ("List Bullet", "List Number"):
        d.styles[name].element.get_or_add_pPr()._remove_numPr()
    d.add_paragraph("{{content}}")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        doc = do_render("a. x\nb. y\n", d, strip_instructions=False)
    assert _num_id(doc.paragraphs[0]) is not None
    assert "lowerLetter" in doc.part.numbering_part.element.xml


def test_lists_in_template_without_numbering_part():
    d = docx.Document()
    rel_id = next(r.rId for r in d.part.rels.values() if r.reltype.endswith("/numbering"))
    d.part.drop_rel(rel_id)
    d.add_paragraph("{{content}}")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        doc = do_render("* a\n\n#. b\n", d, strip_instructions=False)
    assert _num_id(doc.paragraphs[0]) is not None
    assert _num_id(doc.paragraphs[1]) is not None


def test_admonitions_render_as_boxes_and_stay_separate(template):
    doc = do_render(".. note:: n1\n\n.. warning:: w1\n", template)
    boxes = doc.tables[-2:]
    assert boxes[0].cell(0, 0).paragraphs[0].text == "Note"
    assert boxes[1].cell(0, 0).paragraphs[0].text == "Warning"
    assert boxes[0].cell(0, 0).paragraphs[1].text == "n1"


# -- tables ------------------------------------------------------------------------


GRID = """
+-----+-----+
| H1  | H2  |
+=====+=====+
| span both |
+-----+-----+
| a   | b   |
+-----+-----+
"""


def test_grid_table_with_span_and_header(template):
    doc = do_render(GRID, template)
    table = doc.tables[-1]
    assert table.cell(0, 0).text == "H1"
    assert table.cell(1, 0).text == "span both"
    assert table.cell(1, 0)._tc is table.cell(1, 1)._tc
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.cell(0, 0).paragraphs[0].runs[0].bold
    assert table.style.name == "Table Grid"


def test_csv_table_with_title(template):
    rst = '.. csv-table:: Things\n   :header: "A", "B"\n\n   1, 2\n   3, 4\n'
    doc = do_render(rst, template)
    assert "Things" in texts(doc)
    caption = next(p for p in doc.paragraphs if p.text == "Things")
    assert caption.style.name == "Caption"
    table = doc.tables[-1]
    assert [c.text for c in table.rows[2].cells] == ["3", "4"]


def test_var_table(template):
    rst = """
.. var-table:: files
   :columns: name=File:code, sha=SHA-256
   :title: Delivered
"""
    files = [{"name": "a.bin", "sha": "abc"}, {"name": "b.bin", "sha": "def"}]
    doc = do_render(rst, template, variables=dict(VARS, files=files))
    table = doc.tables[-1]
    assert [c.text for c in table.rows[0].cells] == ["File", "SHA-256"]
    assert [c.text for c in table.rows[2].cells] == ["b.bin", "def"]
    assert table.cell(1, 0).paragraphs[0].runs[0].style.name == "HTML Code"


def test_var_table_jira_column_and_empty(template):
    rst = """
.. var-table:: changes
   :columns: key=Issue:jira, summary=Summary

.. var-table:: none
   :columns: key=Issue:jira, summary=Summary
   :empty: No changes.
"""
    changes = [{"key": ["AB-1", "AB-2"], "summary": "fix"}]
    doc = do_render(rst, template, variables=dict(VARS, changes=changes, none=[]))
    links = doc.tables[-2].cell(1, 0)._tc.findall(".//" + qn("w:hyperlink"))
    assert len(links) == 2
    targets = {r.target_ref for r in doc.part.rels.values() if r.is_external}
    assert JIRA + "AB-1" in targets and JIRA + "AB-2" in targets
    empty = doc.tables[-1]
    assert empty.cell(1, 0).text == "No changes."
    assert empty.cell(1, 0)._tc is empty.cell(1, 1)._tc


def test_var_table_errors():
    tpl = simple_template(["{{content}}"])
    with pytest.raises(RstError, match="row 0: variable 'sha' is not defined"):
        render(".. var-table:: files\n   :columns: name=N, sha=S\n", tpl,
               variables={"files": [{"name": "x"}]}, strip_instructions=False)
    tpl.seek(0)
    with pytest.raises(RstError, match="no jira_url"):
        render(".. var-table:: files\n   :columns: name=N:jira\n", tpl,
               variables={"files": []}, strip_instructions=False)


# -- links ---------------------------------------------------------------------------


def test_jira_role(template):
    doc = do_render("Fixed :jira:`AB-12` and :jira:`this one <AB-13>`.", template)
    assert "Fixed AB-12 and this one." in texts(doc)
    targets = {r.target_ref for r in doc.part.rels.values() if r.is_external}
    assert {JIRA + "AB-12", JIRA + "AB-13"} <= targets


def test_jira_url_template():
    doc = do_render(":jira:`AB-1`", simple_template(["{{content}}"]), strip_instructions=False,
                    jira_url="https://j/issues/{key}/view")
    assert "https://j/issues/AB-1/view" in {r.target_ref for r in doc.part.rels.values()}


def test_jira_role_requires_url_and_valid_key():
    tpl = simple_template(["{{content}}"])
    with pytest.raises(RstError, match="no jira_url"):
        render(":jira:`AB-1`", tpl, strip_instructions=False)
    tpl.seek(0)
    with pytest.raises(RstError, match="not a Jira issue key"):
        render(":jira:`nope`", tpl, strip_instructions=False, jira_url=JIRA)


def test_external_link(template):
    doc = do_render("See `docs <https://example.com/x>`_.", template)
    assert "https://example.com/x" in {r.target_ref for r in doc.part.rels.values() if r.is_external}


def test_internal_refs_point_at_bookmarks(template):
    rst = """
.. _install:

Installation
============

text

Other
=====

See :ref:`install`, :ref:`here <install>` and `Other`_.
"""
    doc = do_render(rst, template)
    xml = doc.element.body
    bookmarks = {b.get(qn("w:name")) for b in xml.iter(qn("w:bookmarkStart"))}
    anchors = [h.get(qn("w:anchor")) for h in xml.iter(qn("w:hyperlink"))]
    assert len(anchors) == 3
    assert set(anchors) <= bookmarks
    assert "See Installation, here and Other." in texts(doc)


def test_unknown_ref_raises():
    with pytest.raises(RstError, match="target 'nowhere' not found"):
        render(":ref:`nowhere`", simple_template(["{{content}}"]), strip_instructions=False)


# -- strictness ----------------------------------------------------------------------


def test_unsupported_node_raises(template):
    with pytest.raises(UnsupportedNodeError, match="image"):
        render(".. image:: foo.png\n", template, variables=VARS)


def test_bad_markup_raises():
    with pytest.raises(RstError, match="Unknown interpreted text role"):
        render(":bogus:`x`", simple_template(["{{content}}"]), strip_instructions=False)


def test_missing_style_warns_and_falls_back():
    tpl = simple_template(["{{content}}"])
    with pytest.warns(MissingStyleWarning, match="Plain Text"):
        doc = do_render("::\n\n  code\n", tpl, strip_instructions=False)
    run = doc.paragraphs[0].runs[0]
    assert run.font.name == "Courier New"


def test_source_path_object(tmp_path, template):
    src = tmp_path / "doc.rst"
    src.write_text(":var:`missing`\n")
    with pytest.raises(RstError, match=r"doc\.rst:1:"):
        render(src, template, variables=VARS)
